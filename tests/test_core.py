from io import BytesIO

from docx import Document
from pypdf import PdfWriter
import pytest
from reportlab.pdfgen.canvas import Canvas

from traceclause.matching import match_requirements
from traceclause.parsing import extract_requirements, parse_document


def parse(text):
    return parse_document("test.md", text.encode())


def test_exact_quote_offsets_and_line_numbers():
    doc = parse("# 标题\n\n  系统必须记录操作日志。系统应支持日志导出。\n背景说明")
    clauses = extract_requirements(doc)
    assert len(clauses) == 2
    for clause in clauses:
        block = doc["blocks"][clause["block_id"]]
        assert block["text"][clause["start"]:clause["end"]] == clause["text"]
        assert clause["location"] == "第 3 行"


def test_docx_preserves_paragraph_and_table_order():
    doc = Document()
    doc.add_paragraph("系统必须记录操作日志。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "备份"
    table.cell(0, 1).text = "备份必须保存 30 天。"
    doc.add_paragraph("系统应支持离线录入。")
    output = BytesIO()
    doc.save(output)
    parsed = parse_document("requirements.docx", output.getvalue())
    assert [b["location"] for b in parsed["blocks"]] == ["第 1 段", "表 1 · 第 1 行", "第 2 段"]
    assert len(extract_requirements(parsed)) == 3


def test_pdf_page_location_and_partial_empty_warning():
    output = BytesIO()
    canvas = Canvas(output)
    canvas.drawString(40, 750, "The system must export audit logs.")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    parsed = parse_document("requirements.pdf", output.getvalue())
    assert parsed["blocks"][0]["location"] == "第 1 页"
    assert "第 2 页" in parsed["warnings"][0]
    assert len(extract_requirements(parsed)) == 1


def test_blank_pdf_rejected_instead_of_claiming_no_requirements():
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    with pytest.raises(ValueError, match="OCR"):
        parse_document("scan.pdf", output.getvalue())


@pytest.mark.parametrize("name,data", [("bad.exe", b"abc"), ("bad.docx", b"bad"), ("empty.txt", b""), ("bad.txt", b"\xff\xfe")])
def test_invalid_uploads(name, data):
    with pytest.raises(ValueError):
        parse_document(name, data)


def test_negative_response_flagged_but_not_finally_rejected():
    rows = match_requirements(extract_requirements(parse("系统应支持离线录入。")), parse("本版本不支持离线录入。"))
    assert rows[0]["status"] == "conflict"
    assert rows[0]["review"]["decision"] == "pending"


def test_missing_is_distinct_from_unsupported():
    rows = match_requirements(extract_requirements(parse("系统应提供短信通知能力。")), parse("数据库使用 PostgreSQL。"))
    assert rows[0]["status"] == "missing"
    assert rows[0]["review"]["decision"] == "pending"


def test_numeric_mismatch_is_visible():
    rows = match_requirements(extract_requirements(parse("备份文件必须保存不少于 30 天。")), parse("备份文件保存 7 天。"))
    assert "数值" in rows[0]["reason"]


def test_retrieval_prefers_specific_clause_and_is_deterministic():
    req = extract_requirements(parse("系统应支持 Excel 批量导入设备台账。"))
    res = parse("系统支持用户登录。\n设备台账支持 Excel 批量导入。\n维修工单可以查询。")
    first = match_requirements(req, res)
    assert first[0]["candidates"][0]["block_id"] == 1
    assert first == match_requirements(req, res)
